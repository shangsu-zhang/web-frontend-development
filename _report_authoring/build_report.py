from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\Administrator\Desktop\work\Web前端开发")
DB_ROOT = Path(r"C:\Users\Administrator\Desktop\work\停车场管理数据库项目")
TEMP = ROOT / "_report_authoring"
ASSETS = TEMP / "generated"
TEMPLATE = TEMP / "template-reference.docx"
DB_REPORT = TEMP / "database-report-source.docx"
OUTPUT = ROOT / "停车场管理前端项目文档.docx"
SCREENSHOTS = ROOT / "docs" / "screenshots"


TABLE_NAMES = [
    "administrator", "vehicle_owner", "parking_space", "vehicle",
    "charge_rule", "rental_contract", "parking_record", "payment_record",
]
TABLE_CN = ["管理员", "车主", "车位", "车辆", "收费规则", "月租合同", "停车记录", "缴费记录"]

PAGES = [
    {
        "name": "运营总览", "route": "overview", "image": "overview.png",
        "mapping": "车位状态、在场车辆、近七日已付收入、到期合同",
        "tables": "parking_space、parking_record、payment_record、rental_contract",
        "fields": "space_status、record_status、amount_paid、payment_status、end_date",
        "detail": "展示车位总量与占用率、在场车辆、近七日已支付收入趋势和未来到期合同。指标由统一服务层聚合八表数据，帮助管理员进入具体业务页面。",
    },
    {
        "name": "进出场记录", "route": "parking-records", "image": "parking-records.png",
        "mapping": "车辆与车位关联、入出场时间、停车类型、计费及应收状态",
        "tables": "parking_record、vehicle、parking_space、charge_rule、rental_contract、administrator",
        "fields": "vehicle_id、parking_space_id、parking_type、charge_rule_id、rental_contract_id、entry_time、exit_time、duration_minutes、amount_due、record_status、handled_by_administrator_id",
        "detail": "登记车辆入场并选择可用车位、临停规则或有效月租合同；出场时校验时间并计算临停金额、释放车位，月租停车费用为零。记录保留车辆、车位和办理管理员关联。",
    },
    {
        "name": "缴费记录", "route": "payments", "image": "payments.png",
        "mapping": "支付流水、停车或合同关联、应收实收金额、支付状态",
        "tables": "payment_record、parking_record、rental_contract、administrator",
        "fields": "payment_no、payment_type、parking_record_id、rental_contract_id、amount_due、amount_paid、payment_method、payment_status、paid_at、handled_by_administrator_id、remark",
        "detail": "查看待支付、已支付和已退款流水；可执行模拟收款与退款操作。待支付金额和支付时间状态联动，退款保留原流水金额和时间用于追溯。",
    },
    {
        "name": "查询统计", "route": "reports", "image": "report-01.png",
        "mapping": "11 项 SQL 查询与统计，支持筛选、分页及 CSV 导出",
        "tables": "八表按 SQL 查询说明组合",
        "fields": "3 项单表查询、3 项多表查询、5 项统计；日期、月份、缴费类型、办理管理员筛选",
        "detail": "按数据库查询说明提供当前临停车辆、到期合同、日缴费流水、在场车辆关联、合同收款、分类收入、车位占用、车辆停车、收入总览、管理员办理量和月度每日收入共 11 项结果，并支持分页与 CSV 导出。",
    },
    {
        "name": "车辆档案", "route": "vehicles", "image": "vehicles.png",
        "mapping": "车牌、车型、颜色、车主关联、车辆状态",
        "tables": "vehicle、vehicle_owner",
        "fields": "license_plate、vehicle_type、vehicle_color、owner_id、vehicle_status、registered_at",
        "detail": "支持车辆查询、筛选、排序、详情、新增、编辑和删除未被引用的数据；表单关联车主并校验车牌唯一性及车型、状态取值。",
    },
    {
        "name": "车主档案", "route": "owners", "image": "owners.png",
        "mapping": "姓名、身份证、电话、地址、状态及关联车辆",
        "tables": "vehicle_owner、vehicle",
        "fields": "name、id_number、phone、address、owner_status、created_at",
        "detail": "维护车主实名资料及状态，并展示关联车辆。身份证号、电话按数据库唯一约束校验；车辆记录保留历史关联。",
    },
    {
        "name": "车位管理", "route": "spaces", "image": "spaces.png",
        "mapping": "车位编号、区域、类型、使用状态、位置说明",
        "tables": "parking_space、parking_record",
        "fields": "space_code、area、space_type、space_status、location_note",
        "detail": "按区域浏览车位编号、类型和状态，支持资料维护。车位占用状态由车辆进出场业务更新，禁止把在场车位改成可用。",
    },
    {
        "name": "收费规则", "route": "charge-rules", "image": "charge-rules.png",
        "mapping": "小时费率、最低收费、日封顶、有效期、启停及创建人",
        "tables": "charge_rule、administrator",
        "fields": "rule_name、hourly_rate、minimum_charge、daily_cap、effective_date、expiry_date、rule_status、created_by_administrator_id",
        "detail": "维护临停小时费率、最低收费、每日封顶、有效期和启停状态。金额保留两位小数；NULL 或大于 0 的每日封顶规则与可执行 SQL 一致。",
    },
    {
        "name": "月租合同", "route": "contracts", "image": "contracts.png",
        "mapping": "合同编号、车辆、租期、月租押金、状态及办理人",
        "tables": "rental_contract、vehicle、administrator、payment_record",
        "fields": "contract_no、vehicle_id、start_date、end_date、monthly_rent、deposit、contract_status、signed_at、handled_by_administrator_id、remark",
        "detail": "签订、编辑、续签或终止月租合同，显示车辆与支付情况并检查日期顺序和同车有效期重叠。",
    },
    {
        "name": "管理员", "route": "administrators", "image": "administrators.png",
        "mapping": "员工号、姓名、电话、登录账号、角色、账号状态",
        "tables": "administrator",
        "fields": "employee_no、name、phone、login_account、role、account_status、created_at",
        "detail": "维护管理员资料、岗位角色和账号状态，按员工编号、联系电话、登录账号唯一约束验证输入。",
    },
]


def load_font(size: int, bold: bool = False):
    candidates = (
        [r"C:\Windows\Fonts\msyhbd.ttc", r"C:\Windows\Fonts\simhei.ttf"]
        if bold else [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simsun.ttc"]
    )
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def parse_field_design():
    source = (DB_ROOT / "sql" / "字段设计说明.md").read_text(encoding="utf-8")
    sections = {}
    current = None
    for line in source.splitlines():
        if line.startswith("#"):
            heading = re.match(r"###\s+\d+\.\d+\s+.+?（`([^`]+)`）", line)
            current = heading.group(1) if heading else None
            if current:
                sections[current] = []
            continue
        if current and line.startswith("| "):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if cells and cells[0] not in {"字段名", "---"} and not all(set(cell) <= {"-", ":"} for cell in cells):
                sections[current].append(cells)
    for name in TABLE_NAMES:
        if name not in sections or not sections[name]:
            raise ValueError(f"Missing table definition for {name}")
    return sections


def prepare_schema_rows(sections):
    output = {}
    for sql_name in TABLE_NAMES:
        transformed = []
        for field, data_type, nullable, default, key_index, checks, note in sections[sql_name]:
            match = re.fullmatch(r"([A-Z]+)(?:\(([^)]*)\))?(?:\s+(.*))?", data_type)
            dtype, length, suffix = match.groups() if match else (data_type, "", "")
            if suffix:
                dtype = f"{dtype} {suffix}"
            key_parts = []
            if key_index not in {"—", ""}:
                if sql_name == "parking_record" and field in {"vehicle_id", "parking_space_id"}:
                    columns = "vehicle_id,entry_time" if field == "vehicle_id" else "parking_space_id,entry_time"
                    key_parts.append("INDEX(" + columns + ")")
                elif "联合索引(" in key_index:
                    columns = key_index.split("联合索引(", 1)[1].split(")", 1)[0]
                    key_parts.append("INDEX(" + columns + ")")
                elif "唯一" in key_index:
                    key_parts.append("UNIQUE")
                elif "索引" in key_index:
                    key_parts.append("INDEX")
                elif "主键" in key_index:
                    key_parts.append("PK")
            constraint_parts = ["NOT NULL" if nullable == "否" else "NULL"]
            if default not in {"—", "", "NULL"}:
                constraint_parts.append("AUTO_INCREMENT" if default == "自增" else "DEFAULT " + default)
            if checks not in {"—", ""}:
                fk_match = re.search(r"FK\s*→\s*([A-Za-z_]+\([A-Za-z_]+\))", checks)
                if fk_match:
                    key_parts.append("FK " + fk_match.group(1).replace("(", ".").replace(")", ""))
                    checks = checks.replace(fk_match.group(0), "").strip(" ，,；;：:")
                if checks not in {"—", ""}:
                    if field == "exit_time":
                        checks = "CHECK：NULL 或 >= entry_time；与状态联动"
                    elif field == "duration_minutes":
                        checks = "CHECK：NULL 或 >= 0；与状态联动"
                    elif field == "record_status":
                        checks = "CHECK IN (在场, 已出场, 异常)；与出场时间联动"
                    elif field == "payment_type":
                        checks = "CHECK IN (临停停车费, 月租费)；类型与归属一致"
                    elif field == "parking_record_id" and sql_name == "payment_record":
                        checks = "CHECK：与合同字段二选一"
                    elif field == "rental_contract_id" and sql_name == "payment_record":
                        checks = "CHECK：与停车记录字段二选一"
                    elif "临时停车必须关联收费规则" in checks:
                        checks = "CHECK：临停需收费规则，月租需合同"
                    elif "停车记录与月租合同二选一" in checks:
                        checks = "CHECK：停车记录 / 月租合同二选一"
                    elif "缴费类型与归属对象一致" in checks:
                        checks = "CHECK：缴费类型与归属一致"
                    elif "exit_time、时长为空" in checks or "exit_time、停车时长" in checks:
                        checks = "CHECK：在场时出场时间和时长为空，已出场时非空"
                    elif "待支付为 0" in checks:
                        checks = "CHECK：待支付金额为 0，已支付或退款金额 > 0"
                    elif "待支付为空" in checks:
                        checks = "CHECK：待支付时间为空，已支付或退款时间非空"
                    else:
                        checks = checks.replace("CHECK ∈", "CHECK IN").replace("CHECK 说明：", "CHECK：")
                    constraint_parts.append(checks)
            transformed.append([
                field, dtype, length or "—", "; ".join(constraint_parts),
                "; ".join(key_parts) or "—", note,
            ])
        output[sql_name] = transformed
    return output


def draw_tree_image():
    font = load_font(23)
    rows = [
        (0, "Web前端开发/"),
        (1, "index.html  |  css/style.css"),
        (1, "js/"), (2, "main.js  ui.js  config.js  mock-data.js"),
        (2, "services/  API 与模拟服务适配"),
        (2, "pages/  基础资料、运营、统计页面逻辑"),
        (1, "docs/  Java API、学习基线、页面截图"),
        (2, "screenshots/  运营页与 11 项查询截图"),
        (1, "tests/service.test.cjs  |  README.md  |  AGENTS.md"),
    ]
    image = Image.new("RGB", (1500, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((24, 24, 1476, 476), radius=6, fill="#fbfcfd", outline="#b9c5cc", width=2)
    draw.text((65, 48), "停车场管理前端项目目录", font=load_font(30, True), fill="#162430")
    yy = 112
    for depth, label in rows:
        x = 78 + depth * 44
        if depth:
            draw.text((x - 34, yy), "└─", font=font, fill="#d95035")
        draw.text((x, yy), label, font=font if depth < 2 else load_font(20), fill="#253742")
        yy += 38
    path = ASSETS / "directory-structure.png"
    image.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, font_size=9, bold=False, align=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(font_size)
    run.bold = bold


def update_paragraph_runs(paragraph, text):
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def style_paragraph(paragraph, *, size=11, bold=False, align=None, before=0, after=6, keep=False):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.45
    paragraph.paragraph_format.keep_together = keep
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_text(document, text, *, size=11, bold=False, align=None, before=0, after=6, keep=False):
    paragraph = document.add_paragraph()
    paragraph.add_run(text)
    style_paragraph(paragraph, size=size, bold=bold, align=align, before=before, after=after, keep=keep)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(15 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    for run in paragraph.runs:
        run.font.name = "黑体"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(16 if level == 1 else 14 if level == 2 else 12)
        run.bold = True
    return paragraph


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10)
    return paragraph


def add_page_image(document, image_path, max_width=6.0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.add_run().add_picture(str(image_path), width=Inches(max_width))
    return paragraph


def add_docx_table(document, headers, rows, widths, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        table.columns[index].width = Inches(widths[index])
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_text(cell, header, font_size=font_size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "EAF0F4")
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            set_cell_text(cells[index], value, font_size=font_size)
    return table


def update_cover_table(document):
    table = document.tables[0]
    replacements = {
        2: "待填写（小组全部成员学号）",
        3: "待填写（小组全部成员姓名）",
        4: "待填写",
        5: "2026 年 10 月 22 日",
    }
    for row_index, value in replacements.items():
        cell = table.rows[row_index].cells[1]
        set_cell_text(cell, value, font_size=11)


def set_update_fields(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build_report(sections, schema_images, tree_image):
    document = Document(str(TEMPLATE))
    body = document._element.body
    original_paragraphs = list(document.paragraphs)
    title_p = deepcopy(original_paragraphs[0]._p)
    subtitle_p = deepcopy(original_paragraphs[1]._p)
    metadata_table = deepcopy(document.tables[0]._tbl)

    # Keep the template section and cover structures; replace only its example body.
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    page_break = OxmlElement("w:p")
    run = OxmlElement("w:r")
    break_node = OxmlElement("w:br")
    break_node.set(qn("w:type"), "page")
    run.append(break_node)
    page_break.append(run)
    body.insert(0, title_p)
    body.insert(1, subtitle_p)
    body.insert(2, metadata_table)
    body.insert(3, page_break)

    # Page dimensions and margins come from the retained template.
    document.sections[0].orientation = WD_ORIENT.PORTRAIT
    update_cover_table(document)

    add_heading(document, "目  录", 1)
    contents = [
        "一、项目概述：项目背景、功能简介与任务分配",
        "二、数据库表结构：八张表字段、类型、约束及主外键",
        "三、页面设计与实现：十个前端页面及对应截图",
        "四、页面与数据表的对应关系",
        "五、代码规范与文件结构说明",
        "六、AI 工具交流记录（选做加分项）",
        "七、项目总结",
    ]
    for item in contents:
        add_text(document, item, size=12, after=8)
    add_text(document, "目录页码可在 Word 中打开后，通过“引用 → 目录”插入自动目录并更新。", size=9, after=0)

    document.add_page_break()
    add_heading(document, "一、项目概述", 1)
    add_heading(document, "1.1 项目背景", 2)
    add_text(document, "本项目以《数据库原理与应用》课程的停车场管理数据库设计为基础，继续完成配套前端。数据库项目围绕单一停车场的车主与车辆登记、车位维护、进出场、临时计费、月租合同、缴费及运营查询建立八张核心表。本阶段按课程《Web前端开发技术》项目要求，用原生 HTML、CSS 和 JavaScript 呈现数据库中的字段、关联和业务规则。")
    add_text(document, "页面在浏览器中使用本地模拟数据运行，不连接 MySQL 或 Java 后端；数据访问集中封装，并预留 Java REST 服务适配，后续接入时可配置服务模式与基址。")

    add_heading(document, "1.2 项目功能简介", 2)
    add_text(document, "前端包含运营总览、进出场记录、缴费记录、查询统计、车辆档案、车主档案、车位管理、收费规则、月租合同和管理员十个视图。基础资料页面提供查询、筛选、排序、分页、详情、新增和编辑；业务页面实现入场、出场计费、合同办理、收款与退款；统计页面按数据库查询说明提供三项单表查询、三项多表查询及五项统计，支持筛选、分页和 CSV 导出。")
    add_text(document, "模拟服务对金额、日期、关联、车位占用和合同时间范围进行校验，刷新后保存在浏览器本地存储。界面不是数据库或真实支付系统，支付、退款仅改变演示数据。")

    add_heading(document, "1.3 项目任务分配", 2)
    add_text(document, "实现工作按模块并行推进：基础资料页面、停车与合同缴费业务页面、11 项统计页面、模拟数据与服务适配、浏览器验收和报告整理分别开展。封面要求的小组成员姓名、学号和班级由提交小组依据实际分工填写；本报告不预设成员身份或虚构个人贡献。")

    document.add_page_break()
    add_heading(document, "二、数据库表结构", 1)
    add_text(document, "以下八张表按数据库课程项目的建表 SQL 与字段设计说明整理。每张表先展示一幅结构整理图，再附可编辑字段明细表，便于按图核对并检索字段、类型、长度、约束、主外键和说明。结构图是依据项目文件整理的图示，并非 MySQL Workbench 或 Navicat 的运行截图；字段与约束以可执行 SQL 为准。", size=10)
    for index, (sql_name, cn_name) in enumerate(zip(TABLE_NAMES, TABLE_CN), 1):
        if index > 1:
            document.add_page_break()
        fields = sections[sql_name]
        add_heading(document, f"表 2-{index}  {cn_name}（{sql_name}）", 2)
        add_text(document, f"本表共 {len(fields)} 个字段。主键、唯一键、外键及 CHECK 约束依照数据库定义整理。", size=9)
        add_caption(document, f"图 2-{index}  {cn_name}表结构整理图（依据 SQL 与字段设计说明）")
        add_page_image(document, ASSETS / f"schema-{index:02d}-{sql_name}.png", max_width=5.0)
        add_caption(document, f"表 2-{index}  {cn_name}表字段明细")
        add_docx_table(document,
            ["字段名", "类型", "长度", "约束", "主键 / 外键", "说明"],
            schema_images[sql_name], [1.10, 0.70, 0.42, 1.70, 1.20, 0.88], font_size=7)

    document.add_page_break()
    add_heading(document, "三、页面设计与实现", 1)
    add_text(document, "应用由 index.html 提供统一导航与页面挂载点，各页面通过地址锚点切换；页面代码按基础资料、运营业务和查询统计分组，截图对应本地模拟模式下的实际前端界面。")
    for index, page in enumerate(PAGES, 1):
        if index > 1:
            document.add_page_break()
        add_heading(document, f"3.{index} {page['name']}（index.html#{page['route']}）", 2)
        add_text(document, "对应数据库表：" + page["tables"], size=10)
        add_text(document, "主要字段或查询内容：" + page["fields"], size=9)
        add_text(document, "页面功能说明：" + page["detail"], size=10)
        add_caption(document, f"图 3-{index}  {page['name']}页面截图")
        add_page_image(document, SCREENSHOTS / page["image"], max_width=5.9)

    document.add_page_break()
    add_heading(document, "四、页面与数据表的对应关系", 1)
    add_text(document, "表 4-1 汇总页面入口、所用数据表、主要字段或功能及对应截图编号。运营总览和统计页面会组合多张表；实际 Java 服务接口映射见项目 docs/java-api.md。", size=10)
    add_caption(document, "表 4-1  页面与数据表对应关系表")
    mapping_rows = []
    for index, page in enumerate(PAGES, 1):
        mapping_rows.append([
            str(index), page["name"] + "\n#" + page["route"],
            page["tables"], page["mapping"], f"图 3-{index}",
        ])
    mapping_headers = ["序号", "页面名称（入口）", "对应数据库表", "页面对应字段 / 功能", "截图编号"]
    mapping_widths = [0.38, 1.14, 1.45, 2.48, 0.52]
    add_docx_table(document, mapping_headers, mapping_rows[:5], mapping_widths, font_size=7)
    document.add_page_break()
    add_caption(document, "表 4-1  页面与数据表对应关系表（续）")
    add_docx_table(document, mapping_headers, mapping_rows[5:], mapping_widths, font_size=7)

    document.add_page_break()
    add_heading(document, "五、代码规范与文件结构说明", 1)
    add_heading(document, "5.1 项目文件结构", 2)
    add_text(document, "项目用 index.html 组织导航、总览和视图容器；css/style.css 保存统一样式；js/main.js 负责初始化，js/pages/ 按页面分组实现交互，js/services/ 封装数据读写与 Java API 适配，js/mock-data.js 提供可重置的演示数据。docs/ 保存接口约定、课堂知识映射及页面截图，tests/ 保存 Node 内置测试。")
    add_text(document, "变量和函数使用有意义的英文名称，数据库标识保留 snake_case，中文仅用于界面说明。HTML、CSS、JavaScript 分文件维护；关键模块、页面区块及 JavaScript 函数使用详细中文成对功能注释，例如“功能注释开始 / 功能注释结束”。页面不直接改写存储数据，统一调用 parkingService；当前默认 mock 模式，切换到 Java 模式仅请求预留 REST 接口。")
    add_text(document, "本地运行：直接打开 index.html；测试命令：node --test tests\\service.test.cjs。服务层测试覆盖 CRUD 校验、计费、入场出场、收款退款、浏览器存储恢复及 Java 模式适配。")
    add_caption(document, "图 5-1  项目文件目录结构图")
    add_page_image(document, tree_image, max_width=5.9)

    document.add_page_break()
    add_heading(document, "六、AI 工具交流记录（加分项）", 1)
    add_text(document, "本章按真实项目交流要点作摘要记录，不将摘要伪装成完整对话截图。若课程评分需要逐字交流记录，应由项目成员从实际聊天记录补充。")
    add_heading(document, "6.1 交流记录一", 2)
    add_text(document, "提问摘录：先阅读课程文档和课堂笔记，再基于停车场管理数据库项目搭建前端；以课程要求为最高优先级，课堂笔记作为学习主体，并先生成第一版检查。")
    add_text(document, "处理结果：保留原生 HTML、CSS、JavaScript，将八张表字段、业务操作及 11 项查询映射到页面；补充学习基线说明，区分课程要求与课堂示例。")
    add_heading(document, "6.2 交流记录二", 2)
    add_text(document, "提问摘录：当前阶段只做前端页面、不连接数据库；为以后 Java 前后端连接留接口；代码添加详细成对注释，并优先保证运行能力、性能和美观。")
    add_text(document, "处理结果：把模拟数据持久化和业务校验集中在服务层，配置 Java REST 适配，完成页面路由、窄屏布局和业务流程的浏览器检查。")
    add_heading(document, "6.3 本人对 AI 生成内容的修改与采纳说明", 2)
    add_text(document, "采用了原生脚本结构、模拟服务、REST 接口边界、表单与业务页面组织等实现建议，并根据数据库实际 SQL 修正字段规则和统计口径。保留数据库项目八张表的名称、状态和关联；每日封顶金额按 SQL 采用 NULL 或大于 0。")
    add_text(document, "未采用与当前课程知识范围不匹配的 Vue/React 或额外构建工具；未连接数据库或伪装 Java 服务已上线；未知作者姓名、学号、班级及组员分工仍由小组核实填写。")

    document.add_page_break()
    add_heading(document, "七、项目总结", 1)
    add_text(document, "本项目把 HTML 表格和表单、外部 CSS、选择器、盒子模型、Flex 布局、媒体查询，以及 JavaScript 变量、判断、数组和函数应用到停车场业务页面。以表单录入为例，页面读取控件后调用服务函数校验，模拟服务更新数据，列表和总览再从统一服务读取，形成可解释的前端数据流程。")
    add_text(document, "实现覆盖八表资料与关联、入场占位、出场计费、月租合同、缴费退款及 11 项查询统计。七项服务测试通过；浏览器检查了十个页面路由、统计编号、移动端宽度和控制台错误。页面按本地模拟数据运行，尚未对接真实 Java 后端或 MySQL；下一阶段需依据接口约定完成 Java 服务联调，并补入数据库客户端表结构实机截图及组员身份资料。")

    set_update_fields(document)
    document.core_properties.title = "停车场管理前端项目文档"
    document.core_properties.subject = "Web前端开发技术课程项目报告"
    document.core_properties.author = ""
    document.save(str(OUTPUT))
    return OUTPUT


def main():
    sections = parse_field_design()
    schema_images = prepare_schema_rows(sections)
    tree_image = draw_tree_image()
    output = build_report(sections, schema_images, tree_image)
    print(output)
    print(f"Schema images: {len(schema_images)}; page screenshots: {len(PAGES)}")


if __name__ == "__main__":
    main()
